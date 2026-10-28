import io
import os
import sqlite3
from datetime import datetime
from zoneinfo import ZoneInfo

import streamlit as st
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from PIL import Image, ImageDraw, ImageFont, ImageOps

PASSWORD = os.environ.get("APP_PASSWORD", "tol123")  # GANTI sebelum dipakai tim
TZ = ZoneInfo("Asia/Jakarta")
FOTO_DIR, DB = "foto_watermark", "progres_v2.db"
os.makedirs(FOTO_DIR, exist_ok=True)


def db():
    con = sqlite3.connect(DB)
    con.execute(
        """CREATE TABLE IF NOT EXISTS progres(
        id INTEGER PRIMARY KEY AUTOINCREMENT, waktu TEXT, typical TEXT, zona TEXT,
        area TEXT, sta TEXT, keterangan TEXT, foto TEXT, petugas TEXT)"""
    )
    return con


def dms(v, plus, minus):
    t = round(abs(v) * 3600)
    d, s = divmod(t, 3600)
    m, s = divmod(s, 60)
    return f"{plus if v >= 0 else minus} {d}° {m}' {s}\""


def siapkan_foto(img):
    """Putar sesuai orientasi HP dan kecilkan ukuran. Tidak ada tulisan tambahan."""
    img = ImageOps.exif_transpose(img).convert("RGB")
    img.thumbnail((1280, 1280))
    return img


def buat_excel(rows):
    BIRU, KUNING, NAVY = "0070C0", "FFC000", "1F3864"
    tipis = Side(style="thin")
    garis = Border(left=tipis, right=tipis, top=tipis, bottom=tipis)
    tengah = Alignment(horizontal="center", vertical="center", wrap_text=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Daftar Progres"
    for i, k in enumerate(["Typical", "Zona Kerja", "Area Kerja", "STA Global",
                           "Dokumentasi", "Keterangan"], 1):
        c = ws.cell(row=1, column=i, value=k)
        c.font = Font(name="Times New Roman", size=14, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=NAVY)
        c.alignment, c.border = tengah, garis
    ws.row_dimensions[1].height = 28
    for h, l in zip("ABCDEF", [22, 16, 16, 16, 46, 40]):
        ws.column_dimensions[h].width = l

    def band(r, teks, warna):
        for col in range(1, 7):
            c = ws.cell(row=r, column=col)
            c.fill, c.border = PatternFill("solid", fgColor=warna), garis
        c = ws.cell(row=r, column=1, value=teks)
        c.font = Font(name="Times New Roman", size=14, bold=True)
        ws.row_dimensions[r].height = 24

    grup = {}
    for row in rows:
        grup.setdefault(row["typical"], {}).setdefault(row["zona"], []).append(row)
    r = 1
    for typ, zonas in grup.items():
        r += 1
        band(r, typ, BIRU)
        for zona, items in zonas.items():
            r += 1
            band(r, zona, KUNING)
            for it in items:
                r += 1
                ws.row_dimensions[r].height = 180
                isi = {1: it["typical"], 2: zona, 3: it["area"], 4: it["sta"],
                       6: it["keterangan"]}
                for col in range(1, 7):
                    c = ws.cell(row=r, column=col, value=isi.get(col))
                    c.font = Font(name="Times New Roman", size=12)
                    c.alignment, c.border = tengah, garis
                if os.path.exists(it["foto"]):
                    im = XLImage(it["foto"])
                    k = min(300 / im.width, 225 / im.height)
                    im.width, im.height = int(im.width * k), int(im.height * k)
                    ws.add_image(im, f"E{r}")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


st.set_page_config(page_title="Progres Opname Tol", page_icon="🛣️")
st.title("🛣️ Progres Opname Tol")

if not st.session_state.get("masuk"):
    pw = st.text_input("Password", type="password")
    if st.button("Masuk", type="primary"):
        if pw == PASSWORD:
            st.session_state["masuk"] = True
            st.rerun()
        else:
            st.error("Password salah.")
    st.stop()

tab1, tab2 = st.tabs(["📷 Input", "📋 Daftar & Excel"])

with tab1:
    petugas = st.text_input("Nama QS")
    foto = st.file_uploader("Foto (di HP: pilih 'Ambil foto' atau pilih dari galeri)",
                            type=["jpg", "jpeg", "png"])
    typical = st.text_input("Typical", placeholder="Elevated").strip()
    zona = st.text_input("Zona Kerja", placeholder="Ramp On").strip()
    area = st.text_input("Area Kerja", placeholder="APT").strip()
    c1, c2 = st.columns(2)
    km = c1.number_input("STA Km", min_value=0, value=0, step=1)
    mtr = c2.number_input("STA Meter", min_value=0, max_value=999, value=0, step=10)
    ket = st.text_input("Keterangan", placeholder="Pekerjaan Saluran Drainase").strip()

    if st.button("💾 Simpan", type="primary", use_container_width=True):
        if not (foto and petugas.strip() and typical and zona and area and ket):
            st.error("Foto, nama QS, Typical, Zona, Area, dan Keterangan wajib diisi.")
        else:
            now = datetime.now(TZ)
            hasil = siapkan_foto(Image.open(foto))
            path = os.path.join(FOTO_DIR, now.strftime("%Y%m%d_%H%M%S") + ".jpg")
            hasil.save(path, quality=90)
            con = db()
            con.execute(
                "INSERT INTO progres(waktu,typical,zona,area,sta,keterangan,foto,petugas)"
                " VALUES(?,?,?,?,?,?,?,?)",
                (now.isoformat(), typical, zona, area, f"{km}+{int(mtr):03d}",
                 ket, path, petugas.strip()),
            )
            con.commit()
            con.close()
            st.success("Tersimpan ✅")
            st.image(hasil)


with tab2:
    con = db()
    con.row_factory = sqlite3.Row
    rows = con.execute("SELECT * FROM progres ORDER BY id").fetchall()
    con.close()
    if not rows:
        st.info("Belum ada data.")
    else:
        st.dataframe(
            [{"ID": r["id"], "Waktu": r["waktu"][:16].replace("T", " "),
              "Typical": r["typical"], "Zona": r["zona"], "Area": r["area"],
              "STA": r["sta"], "Keterangan": r["keterangan"], "QS": r["petugas"]}
             for r in rows],
            hide_index=True, use_container_width=True,
        )
        st.download_button(
            "⬇️ Download Excel", data=buat_excel(rows),
            file_name=f"progres_{datetime.now(TZ):%Y%m%d}.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            use_container_width=True,
        )
        with st.expander("Hapus data yang salah input"):
            hapus = st.selectbox("Pilih ID", [r["id"] for r in rows])
            if st.button("🗑️ Hapus"):
                con = db()
                con.execute("DELETE FROM progres WHERE id=?", (hapus,))
                con.commit()
                con.close()
                st.rerun()
